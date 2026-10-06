"""คำนวณ Sales Incentive จากไฟล์ CSV ยอดขาย

ประเภทยอดขาย (ปีที่คำนวณ = Y, เช็คย้อนหลังปี Y-1 ถึง Y-lookback):
  NC   ลูกค้าใหม่              : ลูกค้าไม่มีบิลเลยในช่วงย้อนหลัง
  OCNP ลูกค้าเก่า + สินค้าใหม่ : ลูกค้าเก่า แต่ไม่เคยซื้อ Pline นี้ในช่วงย้อนหลัง
  OCOP ลูกค้าเก่า + สินค้าเก่า : นอกเหนือจากนั้น

Usage: python incentive.py IC3051.txt   (ไฟล์ต้องอยู่ใน data/ — .txt / .csv / .xlsx) [--year 2026] [--config config.yaml]
"""
import argparse
import csv
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import yaml


def load_config(path):
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def parse_date(s):
    s = s.strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    raise ValueError(f"อ่านวันที่ไม่ได้: {s!r}")


def row_date(r, cols):
    if "date" in cols:
        v = r[cols["date"]]
        return v if isinstance(v, datetime) else parse_date(cell_str(v))
    return datetime(int(float(r[cols["year"]])), int(float(r[cols["month"]])), int(float(r[cols["day"]])))


def cell_str(v):
    """แปลงค่าจาก CSV/Excel เป็นข้อความ (เช่น 530011.0 -> '530011')"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_table(path, delimiter=None):
    """อ่าน Raw Text (.txt), CSV หรือ Excel (.xlsx) คืนค่า (header, list ของ dict)"""
    if Path(path).suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [cell_str(h) for h in next(it)]
        rows = [dict(zip(header, r)) for r in it if any(v is not None for v in r)]
        wb.close()
        return header, rows
    with open(path, encoding="utf-8-sig", newline="") as f:
        first = f.readline()
        f.seek(0)
        if not delimiter:
            delimiter = max(["|", "\t", ","], key=first.count)
        reader = csv.DictReader(f, delimiter=delimiter)
        header = [h.strip() for h in reader.fieldnames]
        reader.fieldnames = header
        return header, list(reader)


def load_sales(path, cols, delimiter=None):
    header, raw_rows = read_table(path, delimiter)
    rows = []
    for r in raw_rows:
        rows.append({
            "raw": r,
            "date": row_date(r, cols),
            "salesperson": cell_str(r[cols["salesperson"]]),
            "customer": cell_str(r[cols["customer"]]),
            "pline": cell_str(r[cols["pline"]]),
            "amount": float(cell_str(r[cols["amount"]]).replace(",", "") or 0),
            "margin_raw": cell_str(r.get(cols["margin"])) if "margin" in cols else "",
        })
    return rows, header


def classify(rows, year, lookback):
    """เพิ่ม category ให้ทุกบรรทัดของปี `year`"""
    hist_years = set(range(year - lookback, year))
    hist_customers, hist_cust_pline = set(), set()
    for r in rows:
        if r["date"].year in hist_years:
            hist_customers.add(r["customer"])
            hist_cust_pline.add((r["customer"], r["pline"]))
    out = []
    for r in rows:
        if r["date"].year != year:
            continue
        if r["customer"] not in hist_customers:
            cat = "NC"
        elif (r["customer"], r["pline"]) not in hist_cust_pline:
            cat = "OCNP"
        else:
            cat = "OCOP"
        out.append({**r, "category": cat})
    return out


CATEGORIES = ["OCOP", "OCNP", "NC"]
BELOW = "BELOW LOW"
NO_COMM = "NO COMMISSION"
NEW_PRODUCT = "NEW PRODUCT"


def load_new_products(path):
    """ชีต 'New Product' ใน rates.xlsx: Item No. -> Commission % (Flat rate)"""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    if "New Product" not in wb.sheetnames:
        raise SystemExit(f"❌ {path}: ไม่มีชีต 'New Product' — เพิ่มด้วย: python make_rates_template.py --add-new-product")
    ws = wb["New Product"]
    start = None
    for row in ws.iter_rows():
        if cell_str(row[0].value).lower() == "item no.":
            start = row[0].row + 1
            break
    if start is None:
        raise SystemExit(f"❌ {path}: ชีต New Product ไม่มีหัวคอลัมน์ 'Item No.'")
    items, errors = {}, []
    for r in range(start, ws.max_row + 1):
        item = cell_str(ws.cell(row=r, column=1).value)
        v = ws.cell(row=r, column=2).value
        if not item and (v is None or cell_str(v) == ""):
            continue
        if not item:
            errors.append(f"A{r} (มี % แต่ไม่มี Item No.)")
            continue
        if v is None or cell_str(v) == "":
            errors.append(f"B{r} (Commission % ของ {item})")
            continue
        rate = float(v.strip().rstrip("%")) / 100 if isinstance(v, str) else float(v)
        key = item.upper()
        if key in items:
            errors.append(f"A{r} (Item {item} ซ้ำ)")
        items[key] = rate
    if errors:
        raise SystemExit(f"❌ ชีต New Product ใน {path} ไม่ถูกต้อง:\n  " + "\n  ".join(errors))
    return items


def load_rates(path):
    """อ่าน rates.xlsx
    ตาราง 'TYPE/MARGIN': Commission % ของแต่ละ Type ในแต่ละ Tier
    ตาราง 'MARGIN TYPE' : MIN Margin ของแต่ละ Tier
    คืนค่า list ของ tier เรียงจาก MIN น้อยไปมาก: {name, min, rates{Type: rate}}"""
    import openpyxl
    if not Path(path).exists():
        raise SystemExit(f"ไม่พบไฟล์อัตรา {path} — สร้างด้วย: python make_rates_template.py")
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]

    def find(label):
        for row in ws.iter_rows():
            for c in row:
                if cell_str(c.value).upper() == label:
                    return c.row, c.column
        raise SystemExit(f"❌ {path}: ไม่พบหัวตาราง '{label}' — สร้างใหม่ด้วย make_rates_template.py")

    def num(r, c):
        v = ws.cell(row=r, column=c).value
        if v is None or cell_str(v) == "":
            return None
        if isinstance(v, str):
            return float(v.strip().rstrip("%")) / 100
        return float(v)

    def ref(r, c):
        return f"{ws.cell(row=r, column=c).column_letter}{r}"

    errors = []
    # ตาราง Commission %
    hr, hc = find("TYPE/MARGIN")
    tier_cols = {}
    c = hc + 1
    while cell_str(ws.cell(row=hr, column=c).value):
        tier_cols[cell_str(ws.cell(row=hr, column=c).value).upper()] = c
        c += 1
    type_rows = {}
    r = hr + 1
    while cell_str(ws.cell(row=r, column=hc).value):
        type_rows[cell_str(ws.cell(row=r, column=hc).value).upper()] = r
        r += 1
    missing_types = [t for t in CATEGORIES if t not in type_rows]
    if missing_types or not tier_cols:
        raise SystemExit(f"❌ {path}: ตาราง TYPE/MARGIN ต้องมีแถว {CATEGORIES} และมี Tier อย่างน้อย 1 ช่อง")

    # ตาราง MIN
    mr, mc = find("MARGIN TYPE")
    mins = {}
    r = mr + 1
    while cell_str(ws.cell(row=r, column=mc).value):
        name = cell_str(ws.cell(row=r, column=mc).value).upper()
        v = num(r, mc + 1)
        if v is None:
            errors.append(f"{ref(r, mc + 1)} (MIN ของ {name})")
        mins[name] = v
        r += 1

    tiers = []
    for name, c in tier_cols.items():
        if name not in mins:
            errors.append(f"ตาราง MARGIN TYPE ไม่มี Tier '{name}'")
        rates = {}
        for t in CATEGORIES:
            v = num(type_rows[t], c)
            if v is None:
                errors.append(f"{ref(type_rows[t], c)} ({t} / {name})")
            rates[t] = v
        tiers.append({"name": name, "min": mins.get(name), "rates": rates})
    if errors:
        raise SystemExit(f"❌ กรอก {path} ยังไม่ครบ:\n  " + "\n  ".join(errors))
    tiers.sort(key=lambda t: t["min"])
    for a, b in zip(tiers, tiers[1:]):
        if a["min"] == b["min"]:
            raise SystemExit(f"❌ {path}: MIN ของ {a['name']} และ {b['name']} ซ้ำกัน ({a['min']:.2%})")
    return tiers


def to_margin(raw, amount, mtype):
    """แปลงค่า margin ของแถวเป็นสัดส่วน (0.25 = 25%)"""
    if raw == "":
        return None
    v = float(raw.replace(",", "").rstrip("%"))
    if mtype == "percent":
        return v / 100
    if mtype == "amount":
        return v / amount if amount else None
    return v


def find_tier(m, tiers):
    """Tier ที่ MIN สูงสุดซึ่ง Margin ถึง; ต่ำกว่า MIN ต่ำสุด = None"""
    hit = None
    for t in tiers:
        if m >= t["min"]:
            hit = t
    return hit


def clear_output(out):
    """ลบทุกอย่างในโฟลเดอร์ output ก่อนรัน (กันลบโฟลเดอร์โปรเจกต์หรือโฟลเดอร์แม่โดยไม่ตั้งใจ)"""
    out = out.resolve()
    cwd = Path.cwd().resolve()
    if out == cwd or out in cwd.parents or out == Path.home().resolve():
        raise SystemExit(f"ไม่ลบ {out} — outdir ต้องเป็นโฟลเดอร์ย่อยสำหรับผลลัพธ์เท่านั้น")
    if out.exists():
        for item in out.iterdir():
            shutil.rmtree(item) if item.is_dir() else item.unlink()
    out.mkdir(parents=True, exist_ok=True)


NUMERIC_HINTS = ("price", "qty", "rate", "thb", "amt")
TEXT_COLUMNS = {"currency"}


def is_numeric_col(h):
    """คอลัมน์ตัวเลข (ราคา/จำนวน/อัตรา) — คอลัมน์รหัสอื่น ๆ เก็บเป็นข้อความ กันเลข 0 นำหน้าหาย"""
    hl = h.lower()
    return hl not in TEXT_COLUMNS and any(k in hl for k in NUMERIC_HINTS)


def to_num(v):
    if isinstance(v, (int, float)):
        return v
    t = cell_str(v).replace(",", "")
    try:
        return float(t)
    except ValueError:
        return t


def write_excel(path, year, lookback, month_names, names, monthly, col_total, raw_cols, detail, date_col=None):
    """เขียน 1 ไฟล์ 2 ชีต: Sale Summary และ Raw Data"""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(bold=True, color="FFFFFF")
    total_fill = PatternFill("solid", fgColor="D9E1F2")
    money = "#,##0.00"

    def cell(ws, v, fmt=None, bold=False, fill=None, head=False):
        c = WriteOnlyCell(ws, value=v)
        if head:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if bold:
            c.font = Font(bold=True)
        if fill:
            c.fill = fill
        if fmt:
            c.number_format = fmt
        return c

    # ชีต 1: Sale Summary
    ws = wb.create_sheet("Sale Summary")
    ws.freeze_panes = "B3"
    ws.column_dimensions["A"].width = 12
    for i in range(2, 15):
        ws.column_dimensions[get_column_letter(i)].width = 12
    ws.column_dimensions["N"].width = 14
    ws.append([cell(ws, f"Commission ปี {year} (เช็คลูกค้า/สินค้าย้อนหลังปี {year - lookback}-{year - 1})", bold=True)])
    ws.append([cell(ws, h, head=True) for h in ["Saleman"] + month_names + ["Total"]])
    for n in names:
        ws.append([cell(ws, n)] + [cell(ws, round(v, 2), money) for v in monthly[n]]
                  + [cell(ws, round(sum(monthly[n]), 2), money, bold=True)])
    ws.append([cell(ws, "Total", bold=True, fill=total_fill)]
              + [cell(ws, round(v, 2), money, bold=True, fill=total_fill) for v in col_total]
              + [cell(ws, round(sum(col_total), 2), money, bold=True, fill=total_fill)])

    # ชีต 2: Raw Data
    ws = wb.create_sheet("Raw Data")
    ws.freeze_panes = "A2"
    last = get_column_letter(len(raw_cols))
    ws.auto_filter.ref = f"A1:{last}{len(detail) + 1}"
    for i, h in enumerate(raw_cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 28 if "name" in h.lower() else max(10, min(len(h) + 2, 22))
    ws.append([cell(ws, h, head=True) for h in raw_cols])
    numeric = {h for h in raw_cols if is_numeric_col(h)}
    for d in detail:
        vals = {**d["raw"], "Month": d["date"].month, "Year": d["date"].year, "Type": d["category"],
                "Commission Tier": d["tier"]}
        row = []
        for h in raw_cols:
            if h == date_col:
                row.append(cell(ws, d["date"], "yyyy-mm-dd"))
            elif h == "Calculated Commission %":
                row.append(cell(ws, d["rate"], "0.00%"))
            elif h == "Commission (THB)":
                row.append(cell(ws, d["incentive"], money))
            elif h in ("Month", "Year"):
                row.append(int(vals[h]) if cell_str(vals[h]).isdigit() else vals[h])
            elif h in numeric:
                v = to_num(vals.get(h))
                row.append(cell(ws, v, money) if isinstance(v, float) else v)
            else:
                v = vals.get(h)
                row.append(cell_str(v) if v is not None else None)
        ws.append(row)
    wb.save(path)


DATA_DIR = Path("data")


def in_data(path, what):
    """บังคับให้ไฟล์ input อยู่ในโฟลเดอร์ data/ เท่านั้น (ใส่แค่ชื่อไฟล์ก็ได้)"""
    p = Path(path)
    if not p.is_absolute() and p.parts[:1] != ("data",):
        p = DATA_DIR / p
    data = DATA_DIR.resolve()
    rp = p.resolve()
    if data not in rp.parents:
        raise SystemExit(f"❌ {what} ต้องอยู่ในโฟลเดอร์ data/ เท่านั้น: {path}")
    if not rp.exists():
        raise SystemExit(f"❌ ไม่พบ{what}: {p}")
    return rp


MONTH_NAMES = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def compute(sales_path, rates_path, cfg, year=None):
    """แกนคำนวณ — ใช้ร่วมกันทั้ง CLI และหน้าเว็บ
    คืนค่า dict ผลลัพธ์ หรือ raise SystemExit พร้อมข้อความเมื่อข้อมูลไม่ถูกต้อง"""
    warnings = []
    rows, header = load_sales(sales_path, cfg["columns"], cfg.get("delimiter"))
    excluded = set(cfg.get("exclude_plines") or [])
    rows = [r for r in rows if r["pline"] not in excluded]
    if not rows:
        raise SystemExit("ไม่มียอดขายให้คำนวณ (หลังตัด P-line ที่ยกเว้น)")
    year = year or max(r["date"].year for r in rows)
    lookback = cfg.get("lookback_years", 2)
    years_in_file = {r["date"].year for r in rows}
    missing = [y for y in range(year - lookback, year) if y not in years_in_file]
    if missing:
        warnings.append(f"ไม่มีข้อมูลปี {missing} ในไฟล์ — ลูกค้าอาจถูกนับเป็น NC เกินจริง")

    tiers = load_rates(rates_path)
    new_products = load_new_products(rates_path)
    margin_col = cfg["columns"].get("margin")
    if not margin_col or margin_col not in header:
        raise SystemExit(f"❌ ไม่พบคอลัมน์ Margin '{margin_col}' ในไฟล์ยอดขาย")
    mtype = cfg.get("margin_type", "percent")
    nc_cfg = cfg.get("no_commission") or {}
    skip_items = {str(x).strip().upper() for x in nc_cfg.get("item_nos") or []}
    cost_col = nc_cfg.get("zero_cost_column")
    if cost_col and cost_col not in header:
        raise SystemExit(f"❌ ไม่พบคอลัมน์ต้นทุน '{cost_col}' ในไฟล์ยอดขาย")
    item_col = next((h for h in header if h.strip().lower() == "item no."), None)

    def no_commission(d):
        if item_col and cell_str(d["raw"].get(item_col)).upper() in skip_items:
            return True
        if cost_col:
            v = cell_str(d["raw"].get(cost_col)).replace(",", "")
            return v == "" or float(v) == 0
        return False

    detail = classify(rows, year, lookback)
    if not detail:
        raise SystemExit(f"❌ ไม่มียอดขายของปี {year} ในไฟล์")
    problems = []
    for d in detail:
        if no_commission(d):
            d["tier"], d["rate"], d["incentive"] = NO_COMM, 0.0, 0.0
            continue
        m = to_margin(d["margin_raw"], d["amount"], mtype)
        if m is None:
            problems.append(f"{d['date']:%Y-%m-%d} ลูกค้า {d['customer']} P-line {d['pline']}: ไม่มีค่า Margin")
            continue
        t = find_tier(m, tiers)
        item = cell_str(d["raw"].get(item_col)).upper() if item_col else ""
        if t is None:                      # ต่ำกว่า MIN ต่ำสุด — ไม่ได้ทั้ง Tier และ New Product
            d["tier"], d["rate"] = BELOW, 0.0
        elif item in new_products:         # สินค้าใหม่: Flat rate ตามตาราง
            d["tier"], d["rate"] = NEW_PRODUCT, new_products[item]
        else:
            d["tier"], d["rate"] = t["name"], t["rates"][d["category"]]
        d["incentive"] = round(d["amount"] * d["rate"], 2) + 0.0  # +0.0 กัน -0.0
    if problems:
        raise SystemExit(f"❌ มี {len(problems)} แถวที่หาอัตราไม่ได้ (ตัวอย่าง):\n  " + "\n  ".join(problems[:10]))

    # ---- คอลัมน์ของชีต Raw Data ----
    # ไม่แสดงตัวเลข Margin จริงและต้นทุน — Sale เห็นแค่ Tier
    hidden = {margin_col, *(cfg.get("hide_columns") or [])}
    out_header = [h for h in header if h not in hidden]
    # แยก Month / Year จาก Invoice Date มาไว้ถัดจากคอลัมน์วันที่
    date_col = cfg["columns"].get("date")
    add_my = [c for c in ("Month", "Year") if c not in header]
    if date_col in out_header and add_my:
        i = out_header.index(date_col) + 1
        out_header[i:i] = add_my
    extra = ["Type", "Commission Tier", "Calculated Commission %", "Commission (THB)"]

    # ---- ข้อมูลชีต Sale Summary ----
    monthly = defaultdict(lambda: [0.0] * 12)
    for d in detail:
        monthly[d["salesperson"]][d["date"].month - 1] += d["incentive"]
    names = sorted(monthly)
    col_total = [sum(monthly[n][m] for n in names) for m in range(12)]
    return {"year": year, "lookback": lookback, "warnings": warnings, "names": names,
            "monthly": dict(monthly), "col_total": col_total, "raw_cols": out_header + extra,
            "detail": detail, "date_col": date_col}


def save_excel(res, path):
    """บันทึกผลเป็น Excel (path เป็นชื่อไฟล์ หรือ BytesIO ก็ได้)"""
    write_excel(path, res["year"], res["lookback"], MONTH_NAMES, res["names"], res["monthly"],
                res["col_total"], res["raw_cols"], res["detail"], res["date_col"])


def main():
    p = argparse.ArgumentParser(description="คำนวณ Sales Incentive")
    p.add_argument("sales_csv", help="ไฟล์ยอดขายในโฟลเดอร์ data/ (ใส่แค่ชื่อไฟล์ได้)")
    p.add_argument("--year", type=int, help="ปีที่คำนวณ (ค่าเริ่มต้น = ปีล่าสุดในไฟล์)")
    p.add_argument("--config", default="config.yaml")
    p.add_argument("--outdir", default="output")
    a = p.parse_args()

    cfg = load_config(a.config)
    sales_path = in_data(a.sales_csv, "ไฟล์ยอดขาย")
    rates_path = in_data(cfg.get("rates_file", "rates.xlsx"), "ไฟล์อัตรา (rates.xlsx)")
    res = compute(sales_path, rates_path, cfg, a.year)
    for w in res["warnings"]:
        print(f"⚠️  {w}")

    out = Path(a.outdir)
    clear_output(out)
    out_file = out / f"incentive_{res['year']}.xlsx"
    save_excel(res, out_file)

    year, lookback = res["year"], res["lookback"]
    print(f"Incentive ปี {year} (เช็คย้อนหลังปี {year - lookback}-{year - 1})\n")
    print(f"{'Saleman':<10}" + "".join(f"{m:>9}" for m in MONTH_NAMES) + f"{'Total':>11}")
    for n in res["names"] + ["Total"]:
        vals = res["col_total"] if n == "Total" else res["monthly"][n]
        print(f"{n:<10}" + "".join(f"{v:>9,.0f}" for v in vals) + f"{sum(vals):>11,.2f}")
    print(f"\nบันทึกไฟล์ {out_file}")


if __name__ == "__main__":
    main()
