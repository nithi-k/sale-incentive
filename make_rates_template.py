"""สร้างไฟล์ rates.xlsx (อัตรา Commission ตาม Type x Margin Tier) แบบว่าง ให้กรอกช่องสีเหลือง
Usage: python make_rates_template.py [data/rates.xlsx]
       python make_rates_template.py --add-new-product   (เพิ่มชีต New Product ให้ไฟล์เดิม)
ไม่เขียนทับไฟล์ที่มีอยู่แล้ว (กันอัตราที่กรอกไว้หาย)
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side

TIERS = ["LOW", "MID", "HIGH"]
TYPES = ["OCOP", "OCNP", "NC"]
YELLOW = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def cell(ws, r, c, v=None, bold=False, inp=False, pct=False):
    x = ws.cell(row=r, column=c, value=v)
    x.border = BOX
    x.alignment = Alignment(horizontal="center")
    if bold:
        x.font = Font(bold=True)
    if inp:
        x.fill = YELLOW
        x.protection = Protection(locked=False)
    if pct:
        x.number_format = "0.00%"
    return x


NP_SHEET = "New Product"
NP_ROWS = 200  # จำนวนแถวสีเหลืองเริ่มต้น (เพิ่มต่อท้ายได้)


def add_new_product_sheet(wb):
    """ชีต New Product: Item No. -> Commission % (Flat rate) — ไม่ล็อก เพิ่มแถวได้"""
    ws = wb.create_sheet(NP_SHEET)
    cell(ws, 1, 1, "NEW PRODUCT", bold=True)
    ws.cell(row=1, column=4, value="Item ที่อยู่ในตารางนี้ได้ Commission % ตามตาราง (Flat rate ทุก Type)").font = \
        Font(italic=True, color="666666")
    ws.cell(row=2, column=4, value="ยังต้องผ่าน MIN Margin และไม่ใช่ USDB / ต้นทุน 0").font = Font(italic=True, color="666666")
    cell(ws, 2, 1, "Item No.", bold=True)
    cell(ws, 2, 2, "Commission %", bold=True)
    for r in range(3, 3 + NP_ROWS):
        cell(ws, r, 1, inp=True).alignment = Alignment(horizontal="left")
        cell(ws, r, 2, inp=True, pct=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    ws.freeze_panes = "A3"
    return ws


def build(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rates"
    # ตาราง 1: Commission % ตาม Type x Tier
    cell(ws, 1, 1, "TYPE/MARGIN", bold=True)
    for j, t in enumerate(TIERS, start=2):
        cell(ws, 1, j, t)
    for i, ty in enumerate(TYPES, start=2):
        cell(ws, i, 1, ty)
        for j in range(2, 2 + len(TIERS)):
            cell(ws, i, j, inp=True, pct=True)
    # ตาราง 2: MIN Margin ของแต่ละ Tier
    r0 = 3 + len(TYPES)
    cell(ws, r0, 1, "MARGIN TYPE", bold=True)
    cell(ws, r0, 2, "MIN")
    for i, t in enumerate(TIERS, start=r0 + 1):
        cell(ws, i, 1, t).alignment = Alignment(horizontal="left")
        cell(ws, i, 2, inp=True, pct=True)
    ws["F1"] = "กรอกเฉพาะช่องสีเหลือง"
    ws["F2"] = "Tier = MIN สูงสุดที่ Margin ถึง (Margin ≥ MIN)"
    ws["F3"] = "Margin ต่ำกว่า MIN ของ Tier ต่ำสุด = ไม่ได้ Commission"
    for c in ("F1", "F2", "F3"):
        ws[c].font = Font(italic=True, color="666666")
    ws.column_dimensions["A"].width = 16
    for col in "BCD":
        ws.column_dimensions[col].width = 12
    ws.protection.sheet = True  # แก้ได้เฉพาะช่องสีเหลือง (ไม่มีรหัสผ่าน)
    add_new_product_sheet(wb)
    wb.save(path)


if __name__ == "__main__":
    from pathlib import Path
    if "--add-new-product" in sys.argv:  # เพิ่มชีต New Product ให้ไฟล์เดิม โดยไม่แตะค่าที่กรอกไว้
        from openpyxl import load_workbook
        args = [a for a in sys.argv[1:] if not a.startswith("--")]
        target = args[0] if args else "data/rates.xlsx"
        wb = load_workbook(target)
        if NP_SHEET in wb.sheetnames:
            sys.exit(f"{target} มีชีต {NP_SHEET} อยู่แล้ว")
        add_new_product_sheet(wb)
        wb.save(target)
        print(f"เพิ่มชีต {NP_SHEET} ใน {target} แล้ว")
        sys.exit(0)
    out = sys.argv[1] if len(sys.argv) > 1 else "data/rates.xlsx"
    if Path(out).exists():
        sys.exit(f"มีไฟล์ {out} อยู่แล้ว — ไม่เขียนทับ (ลบหรือเปลี่ยนชื่อไฟล์เดิมก่อน)")
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    build(out)
    print(f"สร้าง {out} แล้ว — กรอกช่องสีเหลืองให้ครบก่อนรัน incentive.py")
